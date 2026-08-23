import docx
import os

md_path = r'C:\Users\cristian andres\.gemini\antigravity-ide\brain\284c50cd-ff24-4210-a745-d2462fc57589\documentacion_historica.md'
output_path = r'C:\Users\cristian andres\OneDrive\Documentos\personal de cristian\privado\Shirokague Devs\IA Shirokage-Aletheia Core\Documentacion_Historica_Aletheia.docx'

with open(md_path, 'r', encoding='utf-8') as f:
    md_text = f.read()

doc = docx.Document()

for line in md_text.split('\n'):
    line = line.strip()
    if not line:
        continue
    if line.startswith('# '):
        doc.add_heading(line[2:], level=0)
    elif line.startswith('## '):
        doc.add_heading(line[3:], level=1)
    elif line.startswith('### '):
        doc.add_heading(line[4:], level=2)
    elif line.startswith('- '):
        doc.add_paragraph(line[2:], style='List Bullet')
    elif line.startswith('>') or line.startswith('[!'):
        # Blockquotes/Alerts
        p = doc.add_paragraph(line)
        p.style = 'Intense Quote'
    else:
        doc.add_paragraph(line)

doc.save(output_path)
print(f'Documento Word actualizado exitosamente en: {output_path}')
